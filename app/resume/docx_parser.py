from pathlib import Path
from app.utils import get_logger

logger = get_logger("docx_parser")


class DOCXParseError(Exception):
    pass


def extract_text_from_docx(docx_path: Path) -> str:
    if not docx_path.exists():
        logger.error("DOCX file not found: %s", docx_path)
        raise DOCXParseError(f"File not found: {docx_path.name}")

    if docx_path.stat().st_size == 0:
        logger.error("DOCX file is empty: %s", docx_path)
        raise DOCXParseError(f"File is empty: {docx_path.name}")

    try:
        from docx import Document
    except ImportError:
        logger.error("python-docx is not installed. Install with: pip install python-docx")
        raise DOCXParseError("DOCX library (python-docx) is not installed")

    try:
        doc = Document(str(docx_path))
    except Exception as exc:
        logger.error("Failed to open DOCX %s: %s", docx_path.name, exc)
        raise DOCXParseError(
            f"Cannot open DOCX: {docx_path.name} — the file may be corrupted or not a valid DOCX"
        ) from exc

    try:
        paragraphs = []
        for para in doc.paragraphs:
            paragraphs.append(para.text)

        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)

        text = "\n".join(paragraphs)

        if not text.strip():
            logger.warning("No text found in DOCX: %s", docx_path.name)

        logger.info(
            "Extracted %d chars from DOCX: %s",
            len(text),
            docx_path.name,
        )
        return text

    except DOCXParseError:
        raise
    except Exception as exc:
        logger.error("Failed to extract text from DOCX %s: %s", docx_path.name, exc)
        raise DOCXParseError(f"Text extraction failed: {exc}") from exc
